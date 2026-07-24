#!/usr/bin/env python3
"""Create a report copy aligned with the verified R001-R006 implementation."""

from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs"
    / "final_report_2026-07-23"
    / "Ng Jing Yong FYP Report - Audited Evaluation Updated.docx"
)
OUTPUT = (
    ROOT
    / "outputs"
    / "final_report_2026-07-23"
    / "Ng Jing Yong FYP Report - Audited Evaluation and System Alignment Updated.docx"
)

OLD = (
    "Fifteen core data and reasoning tests and fifteen backend/API tests passed. "
    "These tests cover source lineage, benchmark gating, controlled action labels, "
    "conflict rules, chronological partitioning, supervised action evidence, "
    "calibration status, recommendation generation, authentication, rule management "
    "and the human-study interface. TypeScript checking and the Vite production build "
    "also completed successfully."
)
NEW = (
    "Fifteen core data and reasoning tests and seventeen backend/API tests passed. "
    "These tests cover source lineage, benchmark gating, controlled action labels, "
    "conflict rules, chronological partitioning, supervised action evidence, "
    "calibration status, recommendation generation, authentication, rule management, "
    "the nine-case counterbalanced human-study interface and the administrator-only "
    "anonymised response export. The Evaluation screen exposes the R001-R006 run "
    "register and preserves pending gates for participant-dependent claims. TypeScript "
    "checking and the Vite production build also completed successfully."
)


def replace_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def main() -> None:
    document = Document(SOURCE)
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == OLD]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one software-verification paragraph, found {len(matches)}."
        )
    replace_paragraph_text(matches[0], NEW)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
