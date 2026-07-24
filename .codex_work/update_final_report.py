#!/usr/bin/env python3
"""Update Chapter 5 and locked-protocol wording with audited evaluation results."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Ng Jing Yong FYP Report After reorganise C3 (Latest).docx"
OUTPUT_DIR = ROOT / "outputs" / "final_report_2026-07-23"
OUTPUT = OUTPUT_DIR / "Ng Jing Yong FYP Report - Audited Evaluation Updated.docx"
METRICS = ROOT / "outputs" / "confirmatory_evaluation" / "metrics.csv"
METADATA = ROOT / "outputs" / "confirmatory_evaluation" / "run_metadata.json"


def replace_paragraph(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def add_before(anchor, text: str, style: str | None = None):
    paragraph = anchor.insert_paragraph_before(text)
    if style:
        paragraph.style = style
    return paragraph


def set_cell_width(cell, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:w"), str(width_dxa))
    width.set(qn("w:type"), "dxa")


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    fill = properties.find(qn("w:shd"))
    if fill is None:
        fill = OxmlElement("w:shd")
        properties.append(fill)
    fill.set(qn("w:fill"), color)


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def set_table_geometry(table, usable_dxa: int, proportions: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(usable_dxa))
    width.set(qn("w:type"), "dxa")
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    widths = [round(usable_dxa * value) for value in proportions]
    widths[-1] += usable_dxa - sum(widths)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])


def add_results_table(document, anchor, rows: list[dict[str, str]], usable_dxa: int):
    headers = ["System", "F1", "Acc.", "Top-2", "Fit", "Conflict", "ECE", "P95 ms"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table._tbl.getparent().remove(table._tbl)
    anchor._p.addprevious(table._tbl)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        shade(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = None
            run.font.size = Pt(7.5)
            run._element.get_or_add_rPr().append(
                OxmlElement("w:color")
            )
            run._element.rPr[-1].set(qn("w:val"), "FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    repeat_header(table.rows[0])
    for row_data in rows:
        cells = table.add_row().cells
        values = [
            row_data["short_name"],
            f"{float(row_data['Macro F1']):.3f}",
            f"{float(row_data['Recommendation Accuracy']):.3f}",
            f"{float(row_data['Top-2 Agreement']):.3f}",
            f"{float(row_data['Scenario-Fit Rate']):.3f}",
            f"{float(row_data['Logical Conflict Rate']):.3f}",
            f"{float(row_data['Expected Calibration Error (ECE)']):.3f}",
            f"{float(row_data['P95 Latency (ms)']):.2f}",
        ]
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[index].paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(7.5)
    set_table_geometry(
        table,
        usable_dxa,
        [.28, .095, .095, .095, .095, .105, .105, .13],
    )
    return table


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document(SOURCE)
    with METRICS.open(newline="", encoding="utf-8") as handle:
        metrics = list(csv.DictReader(handle))
    metadata = json.loads(METADATA.read_text(encoding="utf-8"))

    replacements = {
        "The target benchmark contains at least 500 campaign-period records covering at least three marketing channels and six months of observations. Its unit of analysis is the campaign-period record. Input fields cover campaign context, channel exposure, engagement, cost and outcome variables; derived fields include click-through rate, conversion rate, cost per acquisition, revenue per conversion, ROI band, engagement band and channel saturation. The controlled labels are the business-intent classes Survival, Retention and Growth and the recommendation classes Increase, Maintain, Decrease and Investigate. The final record count, class distribution and excluded records are reported after the reproducible data audit.": (
            "The audited source contains 55,555 campaign-period records covering six "
            "atomic marketing channels and twelve monthly periods. A fixed 125-record "
            "decision-case benchmark was sampled from the source and independently "
            "annotated using Survival, Retention or Growth and Increase, Maintain, "
            "Reduce or Investigate. All sixteen source attributes match the raw layer. "
            "The benchmark was frozen only after third-reviewer adjudication and "
            "completion of the human provenance attestations."
        ),
        "Ground-truth recommendation and intent labels are prepared using documented definitions for the three intents and four actions. Label quality is checked through independent review where more than one annotator is available. Cohen's Kappa measures nominal agreement beyond chance, disagreements are discussed against the written labelling rules, and cases that cannot be resolved are removed before the benchmark is frozen.": (
            "Two independent reviewers labelled all 125 benchmark cases under the "
            "frozen rubric, and a third reviewer adjudicated disagreements. Observed "
            "agreement was 0.904 for intent and 0.784 for action. Cohen's kappa was "
            "0.854 for intent and 0.702 for action, exceeding the predefined 0.60 "
            "action-label threshold. All retained cases were marked Audited before "
            "the benchmark and test split were frozen."
        ),
        "The module applies temperature scaling as a post-processing calibration method. Temperature scaling adjusts predicted probabilities using a validation set without changing the ranking order of recommendations. Expected Calibration Error is then computed by grouping predictions into confidence bins and comparing average confidence with observed correctness in each bin (Guo et al., 2017). Lower ECE indicates better alignment between confidence and actual reliability.": (
            "The implemented module applies temperature scaling using only the "
            "chronological validation partition. Calibration is fitted independently "
            "for the three fixed model seeds and never uses held-out test labels. "
            "Expected Calibration Error is computed with the five predefined bins; "
            "the calibrated full framework reduced mean test ECE from 0.424 in the "
            "without-calibration ablation to 0.226."
        ),
    }
    found = set()
    for paragraph in document.paragraphs:
        current = paragraph.text.strip()
        if current in replacements:
            replace_paragraph(paragraph, replacements[current])
            found.add(current)
        for run in paragraph.runs:
            if "Decrease" in run.text:
                run.text = run.text.replace("Decrease", "Reduce")
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if "Decrease" in run.text:
                            run.text = run.text.replace("Decrease", "Reduce")
    if len(found) != len(replacements):
        raise RuntimeError(f"Could not locate {len(replacements) - len(found)} locked-protocol paragraphs.")

    chapter_five = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == "Chapter 5"
    )
    chapter_six = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == "Chapter 6"
    )
    for paragraph in list(document.paragraphs[chapter_five + 2:chapter_six]):
        delete_paragraph(paragraph)
    anchor = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip() == "Chapter 6"
    )
    add_before(anchor, "Implementation and Testing", "Heading 1")
    add_before(anchor, "5.1 Implementation Status", "Heading 2")
    add_before(
        anchor,
        "The completed prototype implements the five-layer architecture described in "
        "Chapter 4: data management and preprocessing, predictive campaign evidence "
        "and channel attribution, bounded candidate-action generation, weighted "
        "soft-rule and intent-aware ranking, calibrated confidence, structured "
        "explanations, and a React/FastAPI human-in-the-loop interface. The four "
        "recommendation actions are Increase, Maintain, Reduce and Investigate.",
        "Normal Indent",
    )
    add_before(anchor, "5.2 Dataset and Gold Benchmark", "Heading 2")
    add_before(
        anchor,
        "The raw layer contains 55,555 unchanged campaign records and sixteen source "
        "attributes. The processed layer retains the same records and adds "
        "reproducible derived features. The final gold benchmark contains 125 unique "
        "source-matched cases. Reviewer 1, Reviewer 2 and the adjudicator completed "
        "human-only attestations; every accepted row contains valid intent and action "
        "labels and has Annotation_Status set to Audited.",
        "Normal Indent",
    )
    add_before(
        anchor,
        "Inter-rater reliability exceeded the protocol gate: intent observed agreement "
        "was 0.904 with Cohen's kappa of 0.854, while action observed agreement was "
        "0.784 with Cohen's kappa of 0.702. The benchmark was therefore frozen for "
        "technical evaluation.",
        "Normal Indent",
    )
    add_before(anchor, "5.3 Software Verification", "Heading 2")
    add_before(
        anchor,
        "Fifteen core data and reasoning tests and fifteen backend/API tests passed. "
        "These tests cover source lineage, benchmark gating, controlled action labels, "
        "conflict rules, chronological partitioning, supervised action evidence, "
        "calibration status, recommendation generation, authentication, rule "
        "management and the human-study interface. TypeScript checking and the Vite "
        "production build also completed successfully.",
        "Normal Indent",
    )
    add_before(anchor, "5.4 Confirmatory Technical Protocol", "Heading 2")
    add_before(
        anchor,
        "The 125 audited cases were ordered by campaign date and divided into 87 "
        "training cases, 19 validation cases and 19 held-out test cases using the "
        "predefined chronological 70/15/15 split. Temperature scaling was fitted only "
        "on the validation partition. The latest test cases, dated 24 April 2025 to "
        "18 June 2025, were evaluated with identical labels and conflict rules. Model "
        "training used the fixed seeds 42, 52 and 62. Reported intervals are 95% "
        "case-bootstrap intervals based on 2,000 resamples.",
        "Normal Indent",
    )
    add_before(anchor, "5.5 Baseline and Proposed-System Results", "Heading 2")
    add_before(anchor, "Table 5.1 Technical benchmark results", "Normal Indent")
    short_names = {
        "Baseline A - Heuristic Scorecard": "Heuristic",
        "Baseline B - Predictive Analytics Only": "Predictive only",
        "Baseline C - Predictive Analytics with Attribution": "Predictive + attribution",
        "Baseline D - Post-Hoc Rule Overlay": "Post-hoc rules",
        "Proposed System - Full Framework": "Full framework",
        "Proposed Ablation - Without Intent": "Without intent",
        "Proposed Ablation - Without Rules": "Without rules",
        "Proposed Ablation - Without Calibration": "Without calibration",
        "Proposed Ablation - Predictive Backbone": "Predictive backbone",
    }
    for row in metrics:
        row["short_name"] = short_names[row["System"]]
    section = document.sections[0]
    usable_dxa = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
    add_results_table(document, anchor, metrics, usable_dxa)
    add_before(
        anchor,
        "The full framework achieved mean macro F1 of 0.828 and recommendation "
        "accuracy of 0.825, compared with 0.472 and 0.526 for the heuristic baseline. "
        "Its macro-F1 improvement of 0.356 exceeds the predefined 0.05 margin, and its "
        "macro F1 is above the 0.60 continuation threshold. The 95% bootstrap interval "
        "for full-framework macro F1 was 0.603 to 0.969.",
        "Normal Indent",
    )
    add_before(
        anchor,
        "The full framework produced no declared logical conflicts on the held-out "
        "cases and achieved scenario-fit and top-2 agreement of 1.000. The without-rules "
        "ablation achieved higher macro F1 (0.931) but introduced a 0.105 conflict "
        "rate. Removing intent reduced scenario fit to 0.947 and introduced a 0.053 "
        "conflict rate. Calibration reduced ECE from 0.424 to 0.226 without changing "
        "the selected action labels. Full-framework p95 latency was 19.88 ms, well "
        "below the two-second protocol limit.",
        "Normal Indent",
    )
    add_before(anchor, "5.6 Interpretation and Limitations", "Heading 2")
    add_before(
        anchor,
        "The technical results support the initial decision-quality, intent-conditioning "
        "and logical-consistency claims. They also demonstrate a deliberate trade-off: "
        "removing the rule layer improves action agreement on this small test subset but "
        "allows inconsistent recommendations. The full system prioritises conflict-free "
        "decision support while retaining materially higher agreement than the partial "
        "baselines.",
        "Normal Indent",
    )
    add_before(
        anchor,
        "The held-out technical set contains only 19 cases, so the confidence intervals "
        "are wide and the estimates should not be generalised beyond this benchmark. "
        "Actionability proxy, explanation usefulness proxy and trust-alignment proxy are "
        "deterministic technical indicators, not manager ratings. Human actionability, "
        "willingness-to-act, trust-alignment gap and overtrust rate remain pending the "
        "blinded 18-to-24 participant study. No participant-based claim is made here.",
        "Normal Indent",
    )
    add_before(anchor, "5.7 Chapter Summary and Evaluation", "Heading 2")
    add_before(
        anchor,
        "The implemented system passed its software verification, the independently "
        "reviewed benchmark passed the reliability and provenance gates, and the "
        "chronological three-seed technical evaluation met the initial macro-F1 success "
        "criterion. The full framework improved action agreement over the heuristic and "
        "partial analytical baselines, eliminated declared test conflicts, and improved "
        "ECE relative to its uncalibrated ablation. The remaining mandatory activity is "
        "the blinded participant study required for human-reliance conclusions.",
        "Normal Indent",
    )

    document.core_properties.subject = (
        "Audited benchmark and confirmatory technical evaluation update"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
